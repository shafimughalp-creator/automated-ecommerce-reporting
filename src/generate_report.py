"""Write Excel + PNG for the monthly table."""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

DISPLAY = {
    "order_purchase_month": "Month",
    "order_count":          "Orders",
    "orders_with_revenue":  "Orders with revenue",
    "revenue":              "Revenue (BRL)",
    "revenue_per_order":    "Revenue per order (BRL)",
    "coverage_flag":        "Coverage",
}

def write_excel(monthly, out_path):
    display = monthly.rename(columns=DISPLAY)
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Revenue"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, col in enumerate(display.columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for i, row in display.iterrows():
        for j, col in enumerate(display.columns, start=1):
            c = ws.cell(row=i + 2, column=j, value=row[col])
            c.border = border
            if col in ("Revenue (BRL)", "Revenue per order (BRL)"):
                c.number_format = "#,##0.00"
            elif col in ("Orders", "Orders with revenue"):
                c.number_format = "#,##0"
            if col in ("Orders", "Orders with revenue", "Coverage"):
                c.alignment = Alignment(horizontal="center")

    widths = {"Month": 12, "Orders": 10, "Orders with revenue": 20,
              "Revenue (BRL)": 16, "Revenue per order (BRL)": 22,
              "Coverage": 14}
    for j, col in enumerate(display.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths[col]

    ws.freeze_panes = "A2"

    chart = LineChart()
    chart.title = "Monthly Revenue (BRL)"
    chart.y_axis.title = "Revenue (BRL)"
    chart.x_axis.title = "Month"
    chart.height = 9
    chart.width = 22
    chart.style = 2
    data_ref = Reference(ws, min_col=4, min_row=1, max_row=len(display) + 1)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(display) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.series[0].smooth = False
    ws.add_chart(chart, "H2")

    note_row = len(display) + 4
    ws.cell(row=note_row, column=1,
            value=("Coverage: full = complete month; partial_2016 / partial_2018 = "
                   "incomplete raw coverage. 2016-11 is absent. Chart plots all months.")
            ).font = Font(italic=True, size=9)

    ws2 = wb.create_sheet("Data Quality")
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 90
    ws2["A1"] = "Data Quality — Pass 1"
    ws2["A1"].font = Font(bold=True, size=14)
    quality = [
        ("Source", "Brazilian E-Commerce Public Dataset by Olist (historical, public)."),
        ("Disclosure", "Simulated portfolio case study. Monthly batches are SIMULATED."),
        ("Revenue definition", "Payments received on orders the business did not withdraw. NOT cash retained."),
        ("Status exclusions", "canceled + unavailable excluded from revenue."),
        ("Reconciliation", "Order-level and monthly totals independently rebuilt from raw: 0.000000 drift."),
    ]
    for i, (k, v) in enumerate(quality, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out_path)

def write_png(monthly, out_path):
    fig, ax = plt.subplots(figsize=(12, 5))
    full    = monthly[monthly["coverage_flag"] == "full"]
    partial = monthly[monthly["coverage_flag"] != "full"]
    ax.plot(monthly["order_purchase_month"], monthly["revenue"],
            color="#305496", linewidth=1.6, zorder=1)
    ax.scatter(full["order_purchase_month"], full["revenue"],
               color="#305496", s=40, zorder=2, label="Full coverage")
    ax.scatter(partial["order_purchase_month"], partial["revenue"],
               facecolors="white", edgecolors="#C00000", s=55, zorder=3,
               label="Partial coverage")
    ax.set_title("Monthly Revenue (BRL) — Olist dataset, simulated monthly batches")
    ax.set_xlabel("Month")
    ax.set_ylabel("Revenue (BRL)")
    ax.grid(axis="y", linestyle=":", alpha=0.5)
    plt.xticks(rotation=60, ha="right", fontsize=8)
    plt.tight_layout()
    plt.savefig(out_path, dpi=150)
    plt.close(fig)
